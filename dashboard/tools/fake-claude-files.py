"""
A fake Claude Code CLI for the file-sending end-to-end check
(test-file-send.cjs / file-send-check.py).

The agent runtime's Claude Agent SDK starts this as if it were `claude`: it
speaks the same stream-json control protocol on stdin/stdout, so the SDK, its
in-process MCP router and the runtime's send_file server all run for real.
What is fake is only the model: instead of thinking, it does what the prompt
names, the way Claude would through Bash -- it writes the files with the real
libraries (python-docx, openpyxl, reportlab + arabic-reshaper + python-bidi),
then calls send_file through the SDK with control_request/mcp_message.

Prompts it understands (the runtime strips the bot's name first):
  "make the reports"  -> outbox/report.docx, outbox/data.xlsx, outbox/arabic.pdf, each sent
  "make a csv"        -> summary.csv, sent
  "try to escape"     -> asks to send a file outside the folder, a symlink to one,
                         a .sh file and six files (one over the per-reply limit)
  anything else       -> no tool call

Every invocation appends one JSON line to $FAKE_CLAUDE_LOG: cwd, the MCP
servers it was given, whether the system prompt carries the send_file guide,
the prompt, and each tool call with the SDK's answer.
"""
import json
import os
import sys
import time

if len(sys.argv) > 1 and sys.argv[1] in ("-v", "--version"):
    print("2.1.266 (Claude Code)")
    sys.exit(0)

argv = sys.argv[1:]


def arg(name):
    return argv[argv.index(name) + 1] if name in argv else None


mcp_config = json.loads(arg("--mcp-config") or "{}").get("mcpServers", {})
system_prompt = arg("--system-prompt") or ""
LOG = {"cwd": os.getcwd(), "mcp": {k: v.get("type", "stdio") for k, v in mcp_config.items()},
       "allowed": arg("--allowedTools") or arg("--allowed-tools") or "",
       "guide": "SENDING FILES" in system_prompt, "calls": [], "prompt": None}
HAS_TOOL = mcp_config.get("moni_files", {}).get("type") == "sdk"

_req = [0]
_buffer = []


def send(obj):
    sys.stdout.write(json.dumps(obj) + "\n")
    sys.stdout.flush()


def read():
    if _buffer:
        return _buffer.pop(0)
    line = sys.stdin.readline()
    if not line:
        return None
    return json.loads(line)


def answer_control(msg):
    """Reply to a control request from the SDK (initialize, interrupt ...)."""
    req = msg.get("request", {})
    resp = {}
    if req.get("subtype") == "initialize":
        resp = {"commands": [], "output_style": "default", "models": []}
    send({"type": "control_response",
          "response": {"subtype": "success", "request_id": msg["request_id"], "response": resp}})


def mcp(method, params=None):
    """Send an MCP request to the SDK's in-process server, wait for the answer."""
    _req[0] += 1
    rid = "fake_%d" % _req[0]
    message = {"jsonrpc": "2.0", "id": _req[0], "method": method}
    if params is not None:
        message["params"] = params
    send({"type": "control_request", "request_id": rid,
          "request": {"subtype": "mcp_message", "server_name": "moni_files", "message": message}})
    while True:
        msg = read()
        if msg is None:
            raise SystemExit("stdin closed while waiting for " + rid)
        if msg.get("type") == "control_response" and msg["response"].get("request_id") == rid:
            r = msg["response"]
            if r.get("subtype") != "success":
                return {"error": r.get("error")}
            return r["response"]["mcp_response"]
        if msg.get("type") == "control_request":
            answer_control(msg)
        else:
            _buffer.append(msg)


def send_file(path, caption=""):
    res = mcp("tools/call", {"name": "send_file", "arguments": {"path": path, "caption": caption}})
    result = res.get("result") or {}
    text = " ".join(c.get("text", "") for c in result.get("content", []))
    LOG["calls"].append({"path": path, "caption": caption, "ok": not result.get("isError"), "text": text})
    return text


# --- the "model's" work -------------------------------------------------------

def make_reports():
    os.makedirs("outbox", exist_ok=True)
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document()
    d.add_heading("Monthly report", 0)
    d.add_paragraph("Basil sold: 1,200 kg")
    p = d.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("تقرير المبيعات الشهري")
    rtl = OxmlElement("w:rtl")
    run._element.get_or_add_rPr().append(rtl)
    d.save("outbox/report.docx")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Qty kg"])
    ws.append(["Basil", 1200])
    ws.append(["Thyme", 850.5])
    ws2 = wb.create_sheet("عربي")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["المنتج", "الكمية"])
    ws2.append(["ريحان", 1200])
    wb.save("outbox/data.xlsx")

    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table

    pdfmetrics.registerFont(TTFont("Arabic", os.environ.get(
        "FAKE_ARABIC_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")))

    def ar(s):
        return get_display(arabic_reshaper.reshape(s))

    rtl = ParagraphStyle("ar", fontName="Arabic", fontSize=14, leading=20, alignment=TA_RIGHT)
    doc = SimpleDocTemplate("outbox/arabic.pdf", pagesize=A4)
    doc.build([
        Paragraph(ar("تقرير المبيعات الشهري لشركة الجيزة"), rtl),
        Paragraph("Monthly sales report", ParagraphStyle("en", fontName="Arabic")),
        Table([[ar("الكمية"), ar("المنتج")], ["1,200", ar("ريحان")]],
              style=[("FONTNAME", (0, 0), (-1, -1), "Arabic")]),
    ])
    send_file("outbox/report.docx", "Monthly report (Word)")
    send_file("outbox/data.xlsx", "Sales figures (Excel)")
    send_file(os.path.abspath("outbox/arabic.pdf"), "التقرير بالعربي (PDF)")
    return "Sent the Word, Excel and PDF reports."


def make_csv():
    with open("summary.csv", "w", encoding="utf-8") as fh:
        fh.write("product,qty\nBasil,1200\n")
    send_file("summary.csv", "Summary")
    return "Sent summary.csv."


def try_escape():
    outside = os.environ["FAKE_OUTSIDE_FILE"]
    os.makedirs("outbox", exist_ok=True)
    if not os.path.lexists("outbox/link.pdf"):
        os.symlink(outside, "outbox/link.pdf")
    with open("outbox/run.sh", "w") as fh:
        fh.write("echo hi\n")
    send_file(outside, "outside")
    send_file("../../../" + os.path.basename(outside), "dotdot")
    send_file("outbox/link.pdf", "symlink")
    send_file("outbox/run.sh", "script")
    for i in range(6):
        with open("outbox/n%d.txt" % i, "w") as fh:
            fh.write("note %d\n" % i)
        send_file("outbox/n%d.txt" % i, "note %d" % i)
    return "Tried."


def main():
    session = "fake-%d-%d" % (os.getpid(), int(time.time()))
    while True:
        msg = read()
        if msg is None:
            break
        if msg.get("type") == "control_request":
            answer_control(msg)
            continue
        if msg.get("type") != "user":
            continue
        content = msg["message"]["content"]
        prompt = content if isinstance(content, str) else " ".join(
            b.get("text", "") for b in content if isinstance(b, dict))
        LOG["prompt"] = prompt
        if HAS_TOOL:
            tools = mcp("tools/list", {})
            LOG["tools"] = [t["name"] for t in tools.get("result", {}).get("tools", [])]
        low = prompt.lower()
        if not HAS_TOOL and ("make" in low or "escape" in low):
            text = "I have no way to send files here."
        elif "make the reports" in low:
            text = make_reports()
        elif "make a csv" in low:
            text = make_csv()
        elif "try to escape" in low:
            text = try_escape()
        else:
            text = "FAKE-ANSWER " + prompt
        send({"type": "assistant", "message": {"role": "assistant", "model": "fake",
              "content": [{"type": "text", "text": text}]}, "parent_tool_use_id": None,
              "session_id": session})
        send({"type": "result", "subtype": "success", "duration_ms": 5, "duration_api_ms": 1,
              "is_error": False, "num_turns": 1, "session_id": session, "total_cost_usd": 0.0,
              "result": text})
        log = os.environ.get("FAKE_CLAUDE_LOG")
        if log:
            with open(log, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(LOG, ensure_ascii=False) + "\n")


try:
    main()
except Exception as exc:  # make a crash visible to the check
    log = os.environ.get("FAKE_CLAUDE_LOG")
    if log:
        with open(log, "a", encoding="utf-8") as fh:
            fh.write(json.dumps({"crash": repr(exc), "cwd": os.getcwd()}) + "\n")
    raise
